from __future__ import annotations

import asyncio
import csv
from collections.abc import Generator
from datetime import UTC, date, datetime, time
from io import BytesIO, StringIO
from uuid import UUID

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from presensi_api.api.security.dependencies import get_current_user
from presensi_api.api.security.roles import AuthenticatedUser, RoleCode
from presensi_api.db.base import Base
from presensi_api.db.models import (
    AttendanceRecord,
    AttendanceSession,
    Laboratory,
    PracticumSchedule,
    SchoolClass,
    SessionStudent,
    Student,
    User,
)
from presensi_api.db.session import get_db_session
from presensi_api.main import app

ADMIN_ID = UUID("00000000-0000-0000-0000-000000000001")
TEACHER_A_ID = UUID("00000000-0000-0000-0000-000000000002")
TEACHER_B_ID = UUID("00000000-0000-0000-0000-000000000003")
STUDENT_A_ID = UUID("00000000-0000-0000-0000-000000000011")
STUDENT_B_ID = UUID("00000000-0000-0000-0000-000000000012")
STUDENT_C_ID = UUID("00000000-0000-0000-0000-000000000013")
STUDENT_D_ID = UUID("00000000-0000-0000-0000-000000000014")
CLASS_A_ID = UUID("00000000-0000-0000-0000-000000000021")
CLASS_B_ID = UUID("00000000-0000-0000-0000-000000000022")
LAB_A_ID = UUID("00000000-0000-0000-0000-000000000031")
LAB_B_ID = UUID("00000000-0000-0000-0000-000000000032")
SCHEDULE_A_ID = UUID("00000000-0000-0000-0000-000000000041")
SCHEDULE_B_ID = UUID("00000000-0000-0000-0000-000000000042")
SESSION_A_ID = UUID("00000000-0000-0000-0000-000000000051")
SESSION_B_ID = UUID("00000000-0000-0000-0000-000000000052")


@pytest.fixture
def report_database() -> Generator[sessionmaker[Session], None, None]:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    def override_db() -> Generator[Session, None, None]:
        with factory() as session:
            yield session

    app.dependency_overrides[get_db_session] = override_db
    with factory.begin() as db:
        admin = User(
            id=ADMIN_ID,
            email="admin@example.test",
            full_name="Test Admin",
            password_hash="test-hash",
        )
        teacher_a = User(
            id=TEACHER_A_ID,
            email="teacher-a@example.test",
            full_name="Teacher A",
            password_hash="test-hash",
        )
        teacher_b = User(
            id=TEACHER_B_ID,
            email="teacher-b@example.test",
            full_name="Teacher B",
            password_hash="test-hash",
        )
        db.add_all(
            [
                admin,
                teacher_a,
                teacher_b,
                Student(
                    id=STUDENT_A_ID,
                    student_number="S-001",
                    full_name='=HYPERLINK("https://example.test","A")',
                ),
                Student(id=STUDENT_B_ID, student_number="S-002", full_name="Student B"),
                Student(id=STUDENT_C_ID, student_number="S-003", full_name="Student C"),
                Student(id=STUDENT_D_ID, student_number="S-004", full_name="Student D"),
                SchoolClass(
                    id=CLASS_A_ID,
                    code="X-A",
                    name="Class A",
                    grade=10,
                    academic_year="2026-2027",
                ),
                SchoolClass(
                    id=CLASS_B_ID,
                    code="X-B",
                    name="Class B",
                    grade=10,
                    academic_year="2026-2027",
                ),
                Laboratory(id=LAB_A_ID, code="LAB-A", name="Lab A"),
                Laboratory(id=LAB_B_ID, code="LAB-B", name="Lab B"),
                PracticumSchedule(
                    id=SCHEDULE_A_ID,
                    class_id=CLASS_A_ID,
                    laboratory_id=LAB_A_ID,
                    teacher_user_id=TEACHER_A_ID,
                    subject="Chemistry",
                    weekday=1,
                    start_time=time(9),
                    end_time=time(10),
                    timezone_name="Asia/Jakarta",
                    effective_from=date(2026, 1, 1),
                    is_active=True,
                ),
                PracticumSchedule(
                    id=SCHEDULE_B_ID,
                    class_id=CLASS_B_ID,
                    laboratory_id=LAB_B_ID,
                    teacher_user_id=TEACHER_B_ID,
                    subject="Physics",
                    weekday=2,
                    start_time=time(10),
                    end_time=time(11),
                    timezone_name="Asia/Jakarta",
                    effective_from=date(2026, 1, 1),
                    is_active=True,
                ),
                AttendanceSession(
                    id=SESSION_A_ID,
                    practicum_schedule_id=SCHEDULE_A_ID,
                    opened_by_user_id=TEACHER_A_ID,
                    status="active",
                    opened_at=datetime(2026, 9, 15, 2, tzinfo=UTC),
                    grace_period_minutes=15,
                ),
                AttendanceSession(
                    id=SESSION_B_ID,
                    practicum_schedule_id=SCHEDULE_B_ID,
                    opened_by_user_id=TEACHER_B_ID,
                    status="closed",
                    opened_at=datetime(2026, 9, 24, 2, tzinfo=UTC),
                    closed_at=datetime(2026, 9, 24, 3, tzinfo=UTC),
                    grace_period_minutes=15,
                ),
            ]
        )
        db.flush()
        db.add_all(
            [
                SessionStudent(
                    session_id=SESSION_A_ID,
                    student_id=STUDENT_A_ID,
                    student_number_snapshot="S-001",
                    full_name_snapshot='=HYPERLINK("https://example.test","A")',
                ),
                SessionStudent(
                    session_id=SESSION_A_ID,
                    student_id=STUDENT_B_ID,
                    student_number_snapshot="S-002",
                    full_name_snapshot="Student B",
                ),
                SessionStudent(
                    session_id=SESSION_B_ID,
                    student_id=STUDENT_C_ID,
                    student_number_snapshot="S-003",
                    full_name_snapshot="Student C",
                ),
                SessionStudent(
                    session_id=SESSION_B_ID,
                    student_id=STUDENT_D_ID,
                    student_number_snapshot="S-004",
                    full_name_snapshot="Student D",
                ),
                AttendanceRecord(
                    session_id=SESSION_A_ID,
                    student_id=STUDENT_A_ID,
                    status="present",
                    source="manual",
                    recorded_at=datetime(2026, 9, 15, 2, 10, tzinfo=UTC),
                ),
                AttendanceRecord(
                    session_id=SESSION_B_ID,
                    student_id=STUDENT_C_ID,
                    status="late",
                    source="manual",
                    recorded_at=datetime(2026, 9, 24, 2, 30, tzinfo=UTC),
                ),
            ]
        )
    yield factory
    app.dependency_overrides.clear()
    engine.dispose()


def principal(role: RoleCode, user_id: UUID) -> AuthenticatedUser:
    return AuthenticatedUser(
        id=user_id,
        email=f"{role.value.lower()}@example.test",
        full_name=role.value,
        roles=frozenset({role}),
    )


def request(path: str, user: AuthenticatedUser) -> httpx.Response:
    app.dependency_overrides[get_current_user] = lambda: user

    async def send() -> httpx.Response:
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://testserver"
        ) as client:
            return await client.get(path)

    return asyncio.run(send())


def test_report_filters_and_database_pagination(
    report_database: sessionmaker[Session],
) -> None:
    admin = principal(RoleCode.ADMIN, ADMIN_ID)
    base = "/api/v1/reports/attendance/records?starts_on=2026-09-01&ends_on=2026-09-30"

    all_rows = request(base + "&limit=1&offset=1", admin)
    assert all_rows.status_code == 200
    assert all_rows.json()["pagination"] == {"total": 4, "limit": 1, "offset": 1}
    assert len(all_rows.json()["items"]) == 1

    by_student = request(base + f"&student_id={STUDENT_A_ID}", admin).json()
    assert by_student["pagination"]["total"] == 1
    assert by_student["items"][0]["status"] == "present"
    by_identifier = request(base + "&student_number=s-001", admin).json()
    assert by_identifier["pagination"]["total"] == 1
    assert by_identifier["items"][0]["student_id"] == str(STUDENT_A_ID)

    by_class = request(base + f"&class_id={CLASS_B_ID}", admin).json()
    assert by_class["pagination"]["total"] == 2
    by_lab = request(base + f"&laboratory_id={LAB_A_ID}", admin).json()
    assert by_lab["pagination"]["total"] == 2
    by_session = request(base + f"&session_id={SESSION_B_ID}", admin).json()
    assert by_session["pagination"]["total"] == 2
    by_status = request(base + "&status=absent", admin).json()
    assert by_status["pagination"]["total"] == 1
    assert by_status["items"][0]["student_id"] == str(STUDENT_D_ID)

    outside_range = request(
        "/api/v1/reports/attendance/records?starts_on=2026-10-01&ends_on=2026-10-31",
        admin,
    ).json()
    assert outside_range["pagination"]["total"] == 0


def test_teacher_report_is_limited_to_owned_schedules(
    report_database: sessionmaker[Session],
) -> None:
    teacher = principal(RoleCode.TEACHER, TEACHER_A_ID)
    base = "/api/v1/reports/attendance?starts_on=2026-09-01&ends_on=2026-09-30"
    summary = request(base, teacher)
    assert summary.status_code == 200
    assert summary.json()["total_rows"] == 2
    assert summary.json()["present_count"] == 1
    assert summary.json()["not_recorded_count"] == 1

    page = request(
        "/api/v1/reports/attendance/records?starts_on=2026-09-01&ends_on=2026-09-30",
        teacher,
    )
    assert page.status_code == 200
    assert {row["session_id"] for row in page.json()["items"]} == {str(SESSION_A_ID)}

    export = request(
        "/api/v1/reports/attendance/export?"
        "starts_on=2026-09-01&ends_on=2026-09-30&format=csv",
        teacher,
    )
    assert export.status_code == 200
    export_rows = list(csv.reader(StringIO(export.text.lstrip("\ufeff"))))
    session_column = export_rows[0].index("ID sesi")
    assert {row[session_column] for row in export_rows[1:]} == {str(SESSION_A_ID)}


def test_csv_and_xlsx_exports_use_report_fixture_safely(
    report_database: sessionmaker[Session],
) -> None:
    admin = principal(RoleCode.ADMIN, ADMIN_ID)
    common = "starts_on=2026-09-01&ends_on=2026-09-30"
    csv_response = request(
        f"/api/v1/reports/attendance/export?{common}&format=csv", admin
    )
    assert csv_response.status_code == 200, csv_response.text
    assert csv_response.headers["content-type"].startswith("text/csv")
    assert (
        "filename*=UTF-8''attendance-2026-09-01-2026-09-30.csv"
        in csv_response.headers["content-disposition"]
    )
    csv_rows = list(csv.reader(StringIO(csv_response.text.lstrip("\ufeff"))))
    assert csv_rows[0][0] == "Waktu sesi"
    assert len(csv_rows) == 5
    name_column = csv_rows[0].index("Nama siswa")
    assert any(row[name_column].startswith("'=") for row in csv_rows[1:])

    xlsx_response = request(
        f"/api/v1/reports/attendance/export?{common}&format=xlsx", admin
    )
    assert xlsx_response.status_code == 200
    assert xlsx_response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    sheet = load_workbook(BytesIO(xlsx_response.content), read_only=True).active
    assert sheet is not None
    workbook_rows = list(sheet.iter_rows(values_only=True))
    assert len(workbook_rows) == 5
    xlsx_name_column = workbook_rows[0].index("Nama siswa")
    assert any(str(row[xlsx_name_column]).startswith("'=") for row in workbook_rows[1:])
