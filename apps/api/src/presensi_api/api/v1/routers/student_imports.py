from __future__ import annotations

import csv
from dataclasses import dataclass
from io import BytesIO, StringIO
from typing import Annotated, TypedDict
from uuid import UUID
from zipfile import BadZipFile, ZipFile

from fastapi import APIRouter, Depends, File, Form, UploadFile, status
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from presensi_api.api.errors import OPENAPI_ERROR_RESPONSES, ApiProblem
from presensi_api.api.security.dependencies import require_permissions
from presensi_api.api.security.roles import Permission
from presensi_api.api.v1.schemas.common import ErrorDetail
from presensi_api.api.v1.schemas.students import (
    StudentImportCommitResponse,
    StudentImportPreviewResponse,
    StudentImportRowResponse,
)
from presensi_api.db.models import ClassStudent, SchoolClass, Student
from presensi_api.db.session import get_db_session

router = APIRouter(prefix="/students/import", tags=["student-imports"])
DbSession = Annotated[Session, Depends(get_db_session)]
ImportFile = Annotated[UploadFile, File(description="CSV or XLSX student file")]

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10_000
MAX_XLSX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
CSV_MIME_TYPES = {
    "text/csv",
    "application/csv",
    "application/vnd.ms-excel",
    "text/plain",
}
XLSX_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/octet-stream",
}


@dataclass(frozen=True)
class ImportSourceRow:
    row_number: int
    values: tuple[str, ...]


class ParsedImportRow(TypedDict):
    source: ImportSourceRow
    student_number: str
    full_name: str
    class_code: str
    errors: list[str]


@dataclass(frozen=True)
class ValidatedImportRow:
    preview: StudentImportRowResponse
    class_id: UUID | None


@dataclass(frozen=True)
class ValidatedImport:
    response: StudentImportPreviewResponse
    rows: tuple[ValidatedImportRow, ...]


def _problem(status_code: int, code: str, message: str) -> ApiProblem:
    return ApiProblem(status_code, code, message)


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


async def _read_upload(upload: UploadFile) -> tuple[list[str], list[ImportSourceRow]]:
    extension = (
        upload.filename.rsplit(".", 1)[-1].lower()
        if upload.filename and "." in upload.filename
        else ""
    )
    media_type = (upload.content_type or "").split(";", 1)[0].strip().lower()
    if extension == "csv":
        if media_type not in CSV_MIME_TYPES:
            raise _problem(
                415,
                "unsupported_media_type",
                "File CSV harus memakai MIME type text/csv.",
            )
    elif extension == "xlsx":
        if media_type not in XLSX_MIME_TYPES:
            raise _problem(
                415,
                "unsupported_media_type",
                "MIME type tidak sesuai untuk file XLSX.",
            )
    else:
        raise _problem(
            415,
            "unsupported_media_type",
            "Gunakan file dengan ekstensi .csv atau .xlsx.",
        )

    content = await upload.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise _problem(413, "file_too_large", "Ukuran file maksimum adalah 5 MiB.")
    if not content:
        raise _problem(422, "empty_import_file", "File import tidak boleh kosong.")

    try:
        if extension == "csv":
            decoded = content.decode("utf-8-sig")
            table = list(csv.reader(StringIO(decoded)))
        else:
            with ZipFile(BytesIO(content)) as archive:
                expanded_size = sum(item.file_size for item in archive.infolist())
                if expanded_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise _problem(
                        413, "file_too_large", "Isi XLSX melebihi batas pemrosesan."
                    )
                if "xl/workbook.xml" not in archive.namelist():
                    raise _problem(
                        422,
                        "invalid_xlsx_file",
                        "File tidak berisi workbook XLSX yang valid.",
                    )
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            try:
                worksheet = workbook.active
                if worksheet is None:
                    raise _problem(
                        422,
                        "invalid_xlsx_file",
                        "Workbook tidak memiliki worksheet aktif.",
                    )
                table = [
                    [_text(value) for value in row]
                    for row in worksheet.iter_rows(values_only=True)
                ]
            finally:
                workbook.close()
    except UnicodeDecodeError as exc:
        raise _problem(
            422, "invalid_csv_encoding", "File CSV harus memakai UTF-8."
        ) from exc
    except (BadZipFile, csv.Error, InvalidFileException, OSError, ValueError) as exc:
        raise _problem(
            422, "invalid_import_file", "File tidak dapat dibaca sebagai CSV atau XLSX."
        ) from exc

    if not table:
        raise _problem(
            422, "missing_import_header", "File harus memiliki baris header."
        )
    headers = [_text(value) for value in table[0]]
    if not headers or not any(headers):
        raise _problem(
            422, "missing_import_header", "Baris pertama harus berisi nama kolom."
        )
    if len(table) - 1 > MAX_ROWS:
        raise _problem(
            413, "too_many_import_rows", "Import dibatasi maksimum 10.000 baris."
        )

    rows: list[ImportSourceRow] = []
    for row_number, values in enumerate(table[1:], start=2):
        row_text = tuple(_text(value) for value in values)
        if not any(row_text):
            continue
        rows.append(ImportSourceRow(row_number=row_number, values=row_text))
    return headers, rows


def _resolve_mapping(
    headers: list[str], mappings: tuple[str, str, str]
) -> tuple[int, int, int]:
    normalized_headers = [header.casefold() for header in headers]
    if len(set(normalized_headers)) != len(normalized_headers):
        raise _problem(
            422, "duplicate_import_headers", "Nama kolom pada file harus unik."
        )
    if len({name.strip().casefold() for name in mappings}) != 3:
        raise _problem(
            422,
            "invalid_column_mapping",
            "Kolom NIS/NISN, nama, dan kelas harus dipetakan ke header yang berbeda.",
        )
    indexes: list[int] = []
    for mapping in mappings:
        match = next(
            (
                index
                for index, header in enumerate(normalized_headers)
                if header == mapping.strip().casefold()
            ),
            None,
        )
        if match is None:
            raise _problem(
                422,
                "invalid_column_mapping",
                f"Header kolom '{mapping.strip()}' tidak ditemukan pada file.",
            )
        indexes.append(match)
    return indexes[0], indexes[1], indexes[2]


def _validate_rows(
    session: Session,
    headers: list[str],
    source_rows: list[ImportSourceRow],
    mappings: tuple[str, str, str],
) -> ValidatedImport:
    number_index, name_index, class_index = _resolve_mapping(headers, mappings)

    parsed: list[ParsedImportRow] = []
    class_codes: set[str] = set()
    identifiers: set[str] = set()
    for source in source_rows:
        values = source.values

        def value_at(index: int) -> str:
            return values[index].strip() if index < len(values) else ""

        student_number = value_at(number_index).upper()
        full_name = value_at(name_index)
        class_code = value_at(class_index)
        errors: list[str] = []
        if not student_number:
            errors.append("NIS/NISN wajib diisi.")
        elif len(student_number) > 32:
            errors.append("NIS/NISN maksimal 32 karakter.")
        elif student_number.casefold() in identifiers:
            errors.append("NIS/NISN duplikat di dalam file.")
        identifiers.add(student_number.casefold())
        if not full_name:
            errors.append("Nama siswa wajib diisi.")
        elif len(full_name) > 200:
            errors.append("Nama siswa maksimal 200 karakter.")
        if not class_code:
            errors.append("Kode kelas wajib diisi.")
        class_codes.add(class_code.casefold())
        parsed.append(
            {
                "source": source,
                "student_number": student_number,
                "full_name": full_name,
                "class_code": class_code,
                "errors": errors,
            }
        )

    normalized_identifiers = {
        str(row["student_number"]).casefold() for row in parsed if row["student_number"]
    }
    existing_identifiers: set[str] = set()
    if normalized_identifiers:
        existing_identifiers = set(
            session.scalars(
                select(func.lower(Student.student_number)).where(
                    func.lower(Student.student_number).in_(normalized_identifiers)
                )
            ).all()
        )
    for row in parsed:
        identifier = str(row["student_number"]).casefold()
        if identifier and identifier in existing_identifiers:
            row["errors"].append("NIS/NISN sudah terdaftar.")

    matched_classes = (
        session.scalars(
            select(SchoolClass).where(
                func.lower(SchoolClass.code).in_(class_codes),
                SchoolClass.is_active.is_(True),
            )
        ).all()
        if class_codes
        else []
    )
    classes_by_code: dict[str, list[SchoolClass]] = {}
    for school_class in matched_classes:
        classes_by_code.setdefault(school_class.code.casefold(), []).append(
            school_class
        )

    result_rows: list[ValidatedImportRow] = []
    for row in parsed:
        source = row["source"]
        class_code = row["class_code"]
        class_matches = classes_by_code.get(class_code.casefold(), [])
        errors = row["errors"]
        class_id: UUID | None = None
        if not class_matches:
            errors.append(f"Kelas '{class_code}' tidak ditemukan atau nonaktif.")
        elif len(class_matches) > 1:
            errors.append(
                f"Kode kelas '{class_code}' tidak unik "
                "jika huruf besar/kecil diabaikan."
            )
        else:
            class_id = class_matches[0].id
        preview_row = StudentImportRowResponse(
            row_number=source.row_number,
            student_number=str(row["student_number"]),
            full_name=str(row["full_name"]),
            class_code=class_code,
            valid=not errors,
            errors=errors,
        )
        result_rows.append(ValidatedImportRow(preview=preview_row, class_id=class_id))

    valid_count = sum(row.preview.valid for row in result_rows)
    total_rows = len(result_rows)
    response = StudentImportPreviewResponse(
        total_rows=total_rows,
        valid_rows=valid_count,
        invalid_rows=total_rows - valid_count,
        can_commit=total_rows > 0 and valid_count == total_rows,
        rows=[row.preview for row in result_rows],
    )
    return ValidatedImport(response=response, rows=tuple(result_rows))


def _mapping(
    student_number_column: str, full_name_column: str, class_code_column: str
) -> tuple[str, str, str]:
    return student_number_column, full_name_column, class_code_column


async def _validated_import(
    session: Session,
    upload: UploadFile,
    columns: tuple[str, str, str],
) -> ValidatedImport:
    try:
        headers, rows = await _read_upload(upload)
        return _validate_rows(session, headers, rows, columns)
    finally:
        await upload.close()


@router.post(
    "/preview",
    response_model=StudentImportPreviewResponse,
    responses=OPENAPI_ERROR_RESPONSES,
    summary="Preview a student CSV/XLSX import",
    description=(
        "Maps three headers and reports duplicate identifiers, "
        "missing names, and unknown classes before database writes."
    ),
    dependencies=[Depends(require_permissions(Permission.MANAGE_MASTER_DATA))],
)
async def preview_student_import(
    session: DbSession,
    upload: ImportFile,
    student_number_column: Annotated[str, Form(min_length=1, max_length=128)],
    full_name_column: Annotated[str, Form(min_length=1, max_length=128)],
    class_code_column: Annotated[str, Form(min_length=1, max_length=128)],
) -> StudentImportPreviewResponse:
    validated = await _validated_import(
        session,
        upload,
        _mapping(student_number_column, full_name_column, class_code_column),
    )
    return validated.response


@router.post(
    "/commit",
    response_model=StudentImportCommitResponse,
    status_code=status.HTTP_201_CREATED,
    responses=OPENAPI_ERROR_RESPONSES,
    summary="Commit a valid student CSV/XLSX import",
    description=(
        "Revalidates the uploaded file and inserts every row and class link "
        "in one transaction. "
        "Any invalid row or database conflict prevents all writes."
    ),
    dependencies=[Depends(require_permissions(Permission.MANAGE_MASTER_DATA))],
)
async def commit_student_import(
    session: DbSession,
    upload: ImportFile,
    student_number_column: Annotated[str, Form(min_length=1, max_length=128)],
    full_name_column: Annotated[str, Form(min_length=1, max_length=128)],
    class_code_column: Annotated[str, Form(min_length=1, max_length=128)],
) -> StudentImportCommitResponse:
    validated = await _validated_import(
        session,
        upload,
        _mapping(student_number_column, full_name_column, class_code_column),
    )
    if not validated.response.can_commit:
        details = [
            ErrorDetail(
                field=f"rows.{row.preview.row_number}",
                message=message,
                code="invalid_import_row",
            )
            for row in validated.rows
            for message in row.preview.errors
        ]
        raise ApiProblem(
            409,
            "import_rows_invalid",
            "Perbaiki seluruh baris yang tidak valid sebelum commit. "
            "Tidak ada data yang disimpan.",
            details=details,
        )

    try:
        for row in validated.rows:
            student = Student(
                student_number=row.preview.student_number,
                full_name=row.preview.full_name,
            )
            session.add(student)
            session.flush()
            assert row.class_id is not None
            session.add(ClassStudent(class_id=row.class_id, student_id=student.id))
        session.commit()
    except IntegrityError as exc:
        session.rollback()
        raise ApiProblem(
            409,
            "import_conflict",
            "Data berubah sejak preview. Tidak ada data yang disimpan. "
            "Buat preview baru.",
        ) from exc

    return StudentImportCommitResponse(
        imported_count=len(validated.rows),
        class_memberships_created=len(validated.rows),
    )
